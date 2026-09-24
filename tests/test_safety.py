import socket

import httpx
import pytest

from applyapp.config import Settings, write_private_text
from applyapp.documents import create_styled_doc, list_seed_documents
from applyapp.errors import error_message, redact_secrets
from applyapp.google.auth import parse_google_id
from applyapp.posting import (
    UnsafeUrlError,
    _download_html,
    ashby_job_url_from_scripts,
    validate_public_url,
)
from applyapp.queue import fetch_pending_jobs, update_job


def test_public_url_rejects_local_and_obfuscated_addresses():
    blocked = [
        "file:///etc/passwd",
        "http://localhost/jobs",
        "http://127.0.0.1/latest/meta-data/",
        "http://169.254.169.254/latest/meta-data/",
        "http://10.1.2.3/admin",
        "http://2130706433/",
        "http://127.1/",
        "http://0x7f000001/",
        "http://[::ffff:127.0.0.1]/",
        "https://user:secret@jobs.example.com/role",
        "https://metadata.google.internal/computeMetadata/v1/",
    ]
    for url in blocked:
        with pytest.raises(UnsafeUrlError):
            validate_public_url(url)


def test_hostname_resolving_to_a_private_address_is_refused(monkeypatch):
    def fake_resolve(*_args, **_kwargs):
        return [
            (socket.AF_INET, socket.SOCK_STREAM, 0, "", ("1.1.1.1", 0)),
            (socket.AF_INET, socket.SOCK_STREAM, 0, "", ("127.0.0.1", 0)),
        ]

    monkeypatch.setattr("applyapp.posting.socket.getaddrinfo", fake_resolve)
    with pytest.raises(UnsafeUrlError):
        validate_public_url("https://jobs.example.com/role")


def test_redirect_to_loopback_is_refused():
    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.host == "1.1.1.1":
            return httpx.Response(302, headers={"Location": "http://127.0.0.1/latest/meta-data/"})
        raise AssertionError(f"fetched {request.url}")

    client = httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=False)
    with pytest.raises(UnsafeUrlError):
        _download_html(client, "http://1.1.1.1/jobs/1")


def test_embed_script_on_a_private_address_is_not_requested():
    class Client:
        def stream(self, *_args, **_kwargs):
            raise AssertionError("script was fetched")

    html = '<div id="ashby_embed"></div><script src="http://169.254.169.254/latest/meta-data/"></script>'
    page = "https://careers.example.com/jobs?ashby_jid=aaaaaaaa-bbbb-4ccc-8ddd-eeeeeeeeeeee"
    assert ashby_job_url_from_scripts(Client(), page, html) == ""


def test_oversized_page_is_refused(monkeypatch):
    monkeypatch.setattr("applyapp.posting._MAX_RESPONSE_BYTES", 20)

    def handler(_request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, text="x" * 50)

    client = httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=False)
    with pytest.raises(RuntimeError, match="too large"):
        _download_html(client, "http://1.1.1.1/jobs/1")


def test_google_id_rejects_query_text():
    assert parse_google_id("") == ""
    assert parse_google_id("abc' or name contains 'secret") == ""
    assert parse_google_id("https://drive.google.com/drive/folders/1AbC-def_1234567890") == "1AbC-def_1234567890"
    assert parse_google_id("1AbC-def_1234567890") == "1AbC-def_1234567890"


def test_error_text_does_not_keep_api_keys():
    cleaned = redact_secrets("provider said sk-ant-api03-abcdefghijklmnopqrstuvwxyz and Bearer ya29.secret-token")
    assert "sk-ant" not in cleaned
    assert "ya29" not in cleaned
    assert "provider said" in cleaned
    try:
        raise RuntimeError("api_key=sk-openai-abcdefghijklmnopqrstuvwxyz")
    except RuntimeError as exc:
        message = error_message(exc)
    assert "sk-openai" not in message
    assert "[redacted]" in message


def test_queue_writes_formula_text_as_text(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "jobs.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Job Postings", "Company", "Role", "Status"])
    sheet.append(["https://jobs.example.com/a", "Northwind", "Solutions Engineer", "pending"])
    workbook.save(path)
    settings = Settings(job_queue="local", local_jobs_path=str(path))
    job = fetch_pending_jobs(settings)[0]
    update_job(settings, job, company="=HYPERLINK(\"http://evil.example\")", error="@SUM(A1)")
    from openpyxl import load_workbook

    saved = load_workbook(path)
    company = saved.active["B2"]
    assert company.data_type == "s"
    assert not str(company.value).startswith("'")
    assert str(company.value).startswith("=")


def test_env_file_is_private(tmp_path):
    path = tmp_path / ".env"
    write_private_text(path, "ANTHROPIC_API_KEY=sk-ant-example\n")
    assert (path.stat().st_mode & 0o777) == 0o600


def test_seed_symlink_outside_the_folder_is_skipped(tmp_path):
    root = tmp_path / "seeds"
    root.mkdir()
    outside = tmp_path / "secret.md"
    outside.write_text("secret fact", encoding="utf-8")
    (root / "leak.md").symlink_to(outside)
    (root / "Human Writings.md").write_text("voice", encoding="utf-8")
    settings = Settings(
        document_store="local",
        local_seed_dir=str(root),
        local_output_dir=str(tmp_path / "out"),
    )
    files = list_seed_documents(settings)
    assert [item["name"] for item in files] == ["Human Writings.md"]
    assert all("secret fact" not in item["text"] for item in files)


def test_local_output_refuses_a_folder_outside_the_output_root(tmp_path):
    output = tmp_path / "out"
    output.mkdir()
    settings = Settings(
        document_store="local",
        local_seed_dir=str(tmp_path / "seeds"),
        local_output_dir=str(output),
    )
    with pytest.raises(RuntimeError, match="LOCAL_OUTPUT_DIR"):
        create_styled_doc(settings, str(tmp_path), "Resume", [{"kind": "paragraph", "text": "Hi"}])
