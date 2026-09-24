from docx import Document

from applyapp.config import Settings
from applyapp.documents import create_output_folder, create_styled_doc, list_seed_documents, storage_errors


def test_local_store_reads_seeds_and_writes_docx(tmp_path):
    seeds = tmp_path / "seeds" / "Accomplishments"
    seeds.mkdir(parents=True)
    (seeds / "Human Writings.md").write_text("Voice sample.", encoding="utf-8")
    settings = Settings(
        document_store="local",
        local_seed_dir=str(tmp_path / "seeds"),
        local_output_dir=str(tmp_path / "out"),
    )
    files = list_seed_documents(settings)
    assert files[0]["name"] == "Human Writings.md"
    assert files[0]["parents"] == ["Accomplishments"]
    assert "Voice sample" in files[0]["text"]
    assert storage_errors(settings) == []

    folder_id, folder_url = create_output_folder(settings, "2026-09-21 - Northwind - AE")
    url = create_styled_doc(
        settings,
        folder_id,
        "Resume - Northwind - AE",
        [
            {"kind": "name", "text": "Alex Rivera"},
            {"kind": "section", "text": "summary"},
            {"kind": "job_header", "text": "Director          July 2025 – August 2026"},
            {"kind": "bullet", "text": "Shipped the thing"},
        ],
    )
    assert folder_url.startswith("file:")
    assert url.endswith(".docx")
    path = next((tmp_path / "out").rglob("*.docx"))
    document = Document(path)
    name = document.paragraphs[0]
    assert name.text == "Alex Rivera"
    assert name.runs[0].bold
    assert name.runs[0].font.name == "Calibri"
    assert str(name.runs[0].font.color.rgb) == "1A2B44"
    assert document.paragraphs[1].text == "SUMMARY"
    header = document.paragraphs[2].text
    assert header.startswith("Director\t")
    assert "July 2025" in header
    assert "Shipped the thing" in document.paragraphs[3].text


def test_google_store_still_requires_drive_folders():
    settings = Settings(document_store="google", google_seed_folder_id="", google_output_folder_id="")
    assert "GOOGLE_SEED_FOLDER_ID" in storage_errors(settings)
    assert "GOOGLE_OUTPUT_FOLDER_ID" in storage_errors(settings)
