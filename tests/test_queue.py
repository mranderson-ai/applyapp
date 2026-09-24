from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from applyapp.config import Settings
from applyapp.queue import fetch_pending_jobs, needs_google, queue_errors, update_job


def test_local_xlsx_round_trip(tmp_path):
    path = tmp_path / "jobs.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Job Postings", "Company", "Role", "Resume", "Cover Letter", "Status", "Error"])
    sheet.append(["https://jobs.example.com/a", "Northwind", "AE", "", "", "pending", ""])
    sheet.append(["https://jobs.example.com/b", "Contoso", "SE", "", "", "ready_for_review", ""])
    workbook.save(path)

    settings = Settings(job_queue="local", local_jobs_path=str(path))
    jobs = fetch_pending_jobs(settings)
    assert len(jobs) == 1
    assert jobs[0].company == "Northwind"
    update_job(
        settings,
        jobs[0],
        status="ready_for_review",
        company="Northwind Labs",
        resume_doc_url="file:///tmp/resume.docx",
        error="",
    )
    saved = load_workbook(path)
    sheet = saved.active
    headers = [sheet.cell(1, col).value for col in range(1, 9)]
    assert headers[3] == "Posting Text"
    assert sheet["B2"].value == "Northwind Labs"
    assert sheet["G2"].value == "ready_for_review"
    assert sheet["E2"].value == "file:///tmp/resume.docx"
    assert fetch_pending_jobs(settings) == []


def test_missing_workbook_is_created_empty(tmp_path):
    path = tmp_path / "fresh.xlsx"
    settings = Settings(job_queue="local", local_jobs_path=str(path))
    assert fetch_pending_jobs(settings) == []
    saved = load_workbook(path)
    assert saved.active["A1"].value == "Job Postings"
    assert saved.active["D1"].value == "Posting Text"
    assert saved.active.freeze_panes == "A2"
    assert saved.active.data_validations.dataValidation
    assert queue_errors(settings) == []


def test_existing_workbook_gains_posting_text_without_losing_tracking_columns(tmp_path):
    path = tmp_path / "tracker.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "Job Postings",
            "Company",
            "Role",
            "Organization",
            "Level",
            "Resume",
            "Cover Letter",
            "Status",
            "Error",
            "Applied",
            "Notes",
            "Reason",
        ]
    )
    sheet.append(["https://jobs.example.com/a", "Contoso", "Solutions Engineer", "", "", "", "", "pending", "", "", "Yes", ""])
    sheet.column_dimensions["K"].width = 24
    validation = DataValidation(type="list", formula1='"TBD,Yes,No"', allow_blank=True)
    validation.add("K2:K14")
    sheet.add_data_validation(validation)
    workbook.save(path)

    settings = Settings(job_queue="local", local_jobs_path=str(path))
    jobs = fetch_pending_jobs(settings)
    assert len(jobs) == 1
    assert jobs[0].company == "Contoso"
    assert jobs[0].role == "Solutions Engineer"
    assert jobs[0].posting_text == ""

    saved = load_workbook(path)
    headers = [saved.active.cell(1, col).value for col in range(1, 13)]
    assert headers[3] == "Posting Text"
    assert headers[10] == "Applied"
    assert headers[11] == "Notes"
    assert saved.active["B2"].value == "Contoso"
    assert saved.active["L2"].value == "Yes"
    ranges = [str(item.sqref) for item in saved.active.data_validations.dataValidation]
    assert ranges == ["L2:L14"]
    assert saved.active.column_dimensions["L"].width == 24


def test_google_queue_still_requires_a_sheet_and_oauth():
    settings = Settings(job_queue="google", document_store="google", google_sheet_id="")
    assert "GOOGLE_SHEET_ID" in queue_errors(settings)
    assert needs_google(settings) is True


def test_fully_local_setup_does_not_need_google():
    settings = Settings(job_queue="local", document_store="local")
    assert needs_google(settings) is False
