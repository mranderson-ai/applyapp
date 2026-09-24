"""Write a local resume or cover letter as a formatted .docx.

Google Docs is the Drive path. This applies the same `design.STYLES` (Calibri,
navy, 0.75 in margins, section rules, disc bullets) with python-docx. Job dates
use a real right-aligned tab. Docs cannot, so that path still space-pads.
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from applyapp.design import FONT, MARGIN_PT, STYLES
from applyapp.format_blocks import DATE_TAIL

_ALIGN = {
    "CENTER": WD_ALIGN_PARAGRAPH.CENTER,
    "START": WD_ALIGN_PARAGRAPH.LEFT,
    "END": WD_ALIGN_PARAGRAPH.RIGHT,
}
_CONTENT_WIDTH = Inches((612 - 2 * MARGIN_PT) / 72)


def write_styled_docx(path: Path, blocks: list[dict]) -> None:
    """Resume or cover letter. Block kinds match the Docs renderer."""
    document = _blank_document()
    for block in blocks:
        kind = block.get("kind") or "paragraph"
        style = STYLES.get(kind, STYLES["paragraph"])
        text = (block.get("text") or "").strip()
        if style.get("uppercase") and text:
            text = text.upper()
        if kind == "spacer":
            text = ""
        paragraph = document.add_paragraph(style="List Bullet" if style.get("bullet") else None)
        _format_paragraph(paragraph, style)
        if style.get("right_tab"):
            _add_job_header(paragraph, text, style)
        elif text:
            _add_run(paragraph, text, style)
        if style.get("border_bottom"):
            _bottom_rule(paragraph)
    document.save(path)


def write_plain_docx(path: Path, body: str) -> None:
    """QA notes: Calibri body text, same page setup, no resume hierarchy."""
    document = _blank_document()
    style = STYLES["paragraph"]
    for line in body.splitlines() or [""]:
        paragraph = document.add_paragraph()
        _format_paragraph(paragraph, style)
        if line:
            _add_run(paragraph, line, style)
    document.save(path)


def _blank_document() -> Document:
    document = Document()
    section = document.sections[0]
    margin = Inches(MARGIN_PT / 72)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = margin
    section.bottom_margin = margin
    section.left_margin = margin
    section.right_margin = margin
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = _rgb(STYLES["paragraph"]["color"])
    return document


def _format_paragraph(paragraph, style: dict) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = _ALIGN.get(style.get("align") or "START", WD_ALIGN_PARAGRAPH.LEFT)
    fmt.space_before = Pt(style.get("space_before") or 0)
    fmt.space_after = Pt(style.get("space_after") or 0)
    fmt.line_spacing = float(style.get("line_spacing") or 1.08)
    if style.get("bullet"):
        fmt.left_indent = Inches(0.25)
    if style.get("right_tab"):
        fmt.tab_stops.add_tab_stop(_CONTENT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)


def _add_job_header(paragraph, text: str, style: dict) -> None:
    """Title on the left, dates on the right. Collapse the Docs space-padding into a tab."""
    match = DATE_TAIL.search(text)
    if not match:
        _add_run(paragraph, text, style)
        return
    left = text[: match.start()].strip(" |·•-\t")
    right = match.group(1)
    if left:
        _add_run(paragraph, left, style)
    _add_run(paragraph, f"\t{right}", style)


def _add_run(paragraph, text: str, style: dict) -> None:
    run = paragraph.add_run(text)
    run.bold = bool(style.get("bold"))
    run.italic = bool(style.get("italic"))
    run.font.name = FONT
    run.font.size = Pt(style["font_size"])
    run.font.color.rgb = _rgb(style.get("color"))
    fonts = run._element.get_or_add_rPr()
    rfonts = fonts.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        fonts.append(rfonts)
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)


def _bottom_rule(paragraph) -> None:
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")  # 0.75 pt
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1A2B44")
    border.append(bottom)
    paragraph._p.get_or_add_pPr().append(border)


def _rgb(color: dict | None) -> RGBColor:
    color = color or {"red": 0, "green": 0, "blue": 0}
    return RGBColor(round(color["red"] * 255), round(color["green"] * 255), round(color["blue"] * 255))
