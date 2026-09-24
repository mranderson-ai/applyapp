"""Job queue: a Google Sheet, or a local Excel workbook with the same columns.

`JOB_QUEUE=google` is the default. `JOB_QUEUE=local` reads and writes
`LOCAL_JOBS_PATH` (default `jobs.xlsx` in the project root). Header names match
the Sheet, so pending / processing / ready_for_review rules stay in one place.
The workbook is created with a header row, frozen panes, and a filter the first
time it is missing. Cells stay text so Excel does not rewrite URLs.
"""

from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.datavalidation import DataValidation

from applyapp.config import PROJECT_ROOT, Settings
from applyapp.google import sheets as gsheets
from applyapp.models import JobRow

QUEUES = ("google", "local")
LOCAL_HEADERS = [
    "Job Postings",
    "Company",
    "Role",
    "Posting Text",
    "Organization",
    "Level",
    "Resume",
    "Cover Letter",
    "Status",
    "Error",
    "Output",
    "Processed",
]


def fetch_pending_jobs(settings: Settings) -> list[JobRow]:
    if _queue(settings) == "local":
        return gsheets.jobs_from_grid(_read_local(settings))
    return gsheets.fetch_pending_jobs(settings)


def update_job(settings: Settings, job: JobRow, **changes: str) -> JobRow:
    if _queue(settings) == "local":
        return _update_local(settings, job, changes)
    return gsheets.update_job(settings, job, **changes)


def now_iso() -> str:
    return gsheets.now_iso()


def jobs_path(settings: Settings) -> Path:
    """Resolved workbook path. Blank LOCAL_JOBS_PATH means `<project>/jobs.xlsx`."""
    raw = settings.local_jobs_path.strip()
    path = Path(raw).expanduser() if raw else PROJECT_ROOT / "jobs.xlsx"
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")
    return path if path.is_absolute() else PROJECT_ROOT / path


def queue_checks(settings: Settings) -> list[tuple[str, bool, str]]:
    """Doctor lines for the job queue."""
    queue = settings.job_queue.strip().lower() or "google"
    if queue not in QUEUES:
        return [("JOB_QUEUE", False, f"must be google or local, not {queue!r}")]
    if queue == "local":
        path = jobs_path(settings)
        parent_ok = path.parent.is_dir()
        detail = str(path)
        if not path.exists() and parent_ok:
            detail += " (created on first run)"
        return [
            ("JOB_QUEUE", True, "local"),
            ("LOCAL_JOBS_PATH", parent_ok, detail),
        ]
    return [
        ("JOB_QUEUE", True, "google"),
        ("GOOGLE_SHEET_ID", bool(settings.google_sheet_id), ""),
    ]


def queue_errors(settings: Settings) -> list[str]:
    return [name for name, passed, _detail in queue_checks(settings) if not passed]


def needs_google(settings: Settings) -> bool:
    """OAuth is required when the queue or the document store still uses Google."""
    queue = settings.job_queue.strip().lower() or "google"
    store = settings.document_store.strip().lower() or "google"
    return queue != "local" or store != "local"


def _queue(settings: Settings) -> str:
    queue = settings.job_queue.strip().lower() or "google"
    if queue not in QUEUES:
        raise RuntimeError(f"JOB_QUEUE must be google or local, not {queue!r}.")
    return queue


def _read_local(settings: Settings) -> list[list[str]]:
    path = jobs_path(settings)
    if not path.exists():
        _create_workbook(path)
    workbook = load_workbook(path, data_only=False)
    sheet = workbook.active
    if ensure_worksheet_posting_text(sheet):
        workbook.save(path)
    rows: list[list[str]] = []
    for cells in sheet.iter_rows(min_row=1, max_row=sheet.max_row or 1, values_only=True):
        rows.append([_cell_text(value) for value in cells])
    if not rows or not any(rows[0]):
        _create_workbook(path)
        return [LOCAL_HEADERS]
    return rows


def _update_local(settings: Settings, job: JobRow, changes: dict[str, str]) -> JobRow:
    path = jobs_path(settings)
    if not path.exists():
        _create_workbook(path)
    workbook = load_workbook(path)
    sheet = workbook.active
    for field, col in gsheets.field_columns(job).items():
        if field not in changes or col is None:
            continue
        cell = sheet.cell(row=job.sheet_row, column=col + 1, value=changes[field])
        cell.number_format = "@"
        cell.font = Font(name="Calibri", size=11)
    workbook.save(path)
    return job.model_copy(update=changes)


def _create_workbook(path: Path) -> None:
    """Header row a person can filter and type under. Status has a dropdown."""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Jobs"
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1A2B44")
    widths = [52, 22, 28, 48, 22, 16, 46, 46, 22, 40, 46, 22]
    for index, header in enumerate(LOCAL_HEADERS, start=1):
        cell = sheet.cell(1, index, header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = widths[index - 1]
    sheet.row_dimensions[1].height = 22
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(LOCAL_HEADERS))}200"
    sheet["A1"].comment = Comment(
        "Paste a public job URL in this column. Leave Status blank to process the row.",
        "ApplyApp",
        width=240,
        height=40,
    )
    posting_col = LOCAL_HEADERS.index("Posting Text") + 1
    sheet.cell(1, posting_col).comment = Comment(
        gsheets.POSTING_TEXT_NOTE,
        "ApplyApp",
        width=280,
        height=80,
    )
    status_col = get_column_letter(LOCAL_HEADERS.index("Status") + 1)
    validation = DataValidation(
        type="list",
        formula1='"pending,processing,ready_for_review,error"',
        allow_blank=True,
    )
    validation.error = "Use pending, processing, ready_for_review, or error."
    validation.errorTitle = "Status"
    validation.add(f"{status_col}2:{status_col}200")
    sheet.add_data_validation(validation)
    for row in range(2, 201):
        for col in range(1, len(LOCAL_HEADERS) + 1):
            sheet.cell(row, col).number_format = "@"
    workbook.save(path)


def ensure_worksheet_posting_text(sheet) -> bool:
    """Insert Posting Text after Role on a workbook that was created before the column existed.

    Cell values move with the column. Dropdowns, merged ranges, and column widths
    are shifted with them so tracking columns such as Applied stay attached to
    their data.
    """
    last_col = sheet.max_column or 1
    headers = [_cell_text(sheet.cell(1, col).value) for col in range(1, last_col + 1)]
    while headers and not headers[-1]:
        headers.pop()
    insert_at = gsheets.posting_text_insert_at(headers)
    if insert_at is None:
        return False
    column = insert_at + 1
    merges = list(sheet.merged_cells.ranges)
    for cell_range in merges:
        sheet.unmerge_cells(str(cell_range))
    sheet.insert_cols(column)
    for cell_range in merges:
        min_col, max_col = cell_range.min_col, cell_range.max_col
        if min_col >= column:
            min_col += 1
            max_col += 1
        elif max_col >= column:
            max_col += 1
        sheet.merge_cells(
            start_row=cell_range.min_row,
            start_column=min_col,
            end_row=cell_range.max_row,
            end_column=max_col,
        )
    _shift_data_validations(sheet, column)
    _shift_column_widths(sheet, column)
    _write_posting_text_header(sheet, column)
    return True


def _shift_data_validations(sheet, column: int) -> None:
    for validation in sheet.data_validations.dataValidation:
        shifted: list[str] = []
        for cell_range in validation.sqref.ranges:
            min_col = cell_range.min_col
            max_col = cell_range.max_col
            if min_col >= column:
                min_col += 1
                max_col += 1
            elif max_col >= column:
                max_col += 1
            shifted.append(
                str(
                    CellRange(
                        min_col=min_col,
                        min_row=cell_range.min_row,
                        max_col=max_col,
                        max_row=cell_range.max_row,
                    )
                )
            )
        validation.sqref = " ".join(shifted)


def _shift_column_widths(sheet, column: int) -> None:
    widths: dict[int, float] = {}
    for letter, dim in list(sheet.column_dimensions.items()):
        if not letter or dim.width is None:
            continue
        widths[column_index_from_string(letter)] = dim.width
    for index in sorted((index for index in widths if index >= column), reverse=True):
        sheet.column_dimensions[get_column_letter(index + 1)].width = widths[index]
        del sheet.column_dimensions[get_column_letter(index)]
    sheet.column_dimensions[get_column_letter(column)].width = 48


def _write_posting_text_header(sheet, column: int) -> None:
    source = sheet.cell(1, column - 1) if column > 1 else None
    cell = sheet.cell(1, column, gsheets.POSTING_TEXT_HEADER)
    if source is not None and source.has_style:
        cell.font = copy(source.font)
        cell.fill = copy(source.fill)
        cell.alignment = copy(source.alignment)
        cell.border = copy(source.border)
    cell.comment = Comment(gsheets.POSTING_TEXT_NOTE, "ApplyApp", width=280, height=80)
    for row in range(2, max(sheet.max_row or 1, 2) + 1):
        sheet.cell(row, column).number_format = "@"


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%dT%H:%M:%SZ")
    return str(value).strip()
